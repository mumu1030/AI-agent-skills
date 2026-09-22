#!/usr/bin/env python3
"""
从 Excel 测试用例表读取结构化用例列表，输出 JSON 供 opentest-case-review 评审使用。

用法：
    python read_cases_excel.py --xlsx path/to/cases.xlsx
    python read_cases_excel.py --xlsx path/to/cases.xlsx --sheet 测试用例
"""
import argparse
import json
import sys
from typing import Any, Dict, List, Optional


# 标准字段 -> 可能的表头别名
COLUMN_ALIASES: Dict[str, List[str]] = {
    "case_id": ["用例编号", "编号", "ID", "用例ID", "case id", "Case ID"],
    "name": ["用例名称", "用例标题", "标题", "名称", "测试用例", "case name"],
    "feature": ["功能模块", "模块", "所属模块", "功能", "feature"],
    "type": ["类型", "用例类型", "测试类型", "type"],
    "preconditions": ["前置条件", "前置", "预置条件", "preconditions"],
    "steps": ["测试步骤", "步骤", "操作步骤", "执行步骤", "steps"],
    "expected": ["预期结果", "预期", "期望结果", "expected"],
    "priority": ["优先级", "优先级别", "priority"],
    "test_data": ["测试数据", "数据", "输入数据", "test data"],
    "requirement_ref": ["需求编号", "需求ID", "关联需求", "需求", "requirement"],
    "execution_result": ["执行结果", "结果", "执行状态"],
}


def _normalize_header(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _build_column_map(headers: List[Any]) -> Dict[str, int]:
    """将第一行表头映射到标准字段名 -> 列索引（0-based）。"""
    col_map: Dict[str, int] = {}
    normalized_headers = [_normalize_header(h) for h in headers]

    for field, aliases in COLUMN_ALIASES.items():
        if field in col_map:
            continue
        for idx, header in enumerate(normalized_headers):
            if not header:
                continue
            for alias in aliases:
                if header == alias or header.lower() == alias.lower():
                    col_map[field] = idx
                    break
            if field in col_map:
                break

    return col_map


def _cell_str(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value == int(value):
        return str(int(value))
    return str(value).strip()


def read_cases_from_xlsx(xlsx_path: str, sheet_name: Optional[str] = None) -> List[Dict[str, Any]]:
    try:
        from openpyxl import load_workbook
    except ImportError:
        raise RuntimeError("需要安装 openpyxl: pip install openpyxl")

    wb = load_workbook(xlsx_path, read_only=True, data_only=True)
    if sheet_name and sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
    elif "测试用例" in wb.sheetnames:
        ws = wb["测试用例"]
    else:
        ws = wb.active

    headers: List[Any] = []
    for col in range(1, ws.max_column + 1):
        headers.append(ws.cell(row=1, column=col).value)

    col_map = _build_column_map(headers)
    cases: List[Dict[str, Any]] = []

    def get_field(row: int, field: str) -> str:
        if field not in col_map:
            return ""
        col_idx = col_map[field] + 1
        return _cell_str(ws.cell(row=row, column=col_idx).value)

    for row in range(2, ws.max_row + 1):
        name = get_field(row, "name")
        steps = get_field(row, "steps")
        expected = get_field(row, "expected")
        case_id = get_field(row, "case_id")

        # 跳过空行
        if not any([name, steps, expected, case_id]):
            continue

        case: Dict[str, Any] = {
            "row": row,
            "case_id": case_id or f"ROW-{row}",
            "name": name,
            "feature": get_field(row, "feature"),
            "type": get_field(row, "type"),
            "preconditions": get_field(row, "preconditions"),
            "steps": steps,
            "expected": expected,
            "priority": get_field(row, "priority"),
            "test_data": get_field(row, "test_data"),
            "requirement_ref": get_field(row, "requirement_ref"),
            "execution_result": get_field(row, "execution_result"),
            "source": "excel",
            "source_file": xlsx_path,
        }
        cases.append(case)

    wb.close()
    return cases


def main() -> None:
    parser = argparse.ArgumentParser(description="Read test cases from Excel for case review")
    parser.add_argument("--xlsx", "-x", required=True, help="Path to .xlsx test case file")
    parser.add_argument("--sheet", "-s", default=None, help="Sheet name (default: 测试用例 or active)")
    args = parser.parse_args()

    try:
        cases = read_cases_from_xlsx(args.xlsx, args.sheet)
        out = {
            "source": args.xlsx,
            "sheet": args.sheet or "测试用例",
            "count": len(cases),
            "cases": cases,
        }
        print(json.dumps(out, ensure_ascii=False, indent=2))
    except Exception as e:
        print(json.dumps({"error": str(e), "path": args.xlsx}, ensure_ascii=False), file=sys.stderr)
        sys.exit(1)


if __name__ == "__main__":
    main()
