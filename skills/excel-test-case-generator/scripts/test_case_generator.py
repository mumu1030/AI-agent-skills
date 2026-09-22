#!/usr/bin/env python3
"""
Test Case Generator - Creates formatted test case Excel files

This script generates well-formatted Excel files for test cases with standard
structure including test case ID, name, type, preconditions, steps, expected
results, priority, and execution results.
"""

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from typing import List, Dict, Optional, Tuple
from pathlib import Path
import json
try:
    import yaml
except ImportError:
    yaml = None


class TestCaseGenerator:
    """Generator for test case Excel files."""
    
    # Standard column headers
    HEADERS = [
        "用例编号",
        "用例名称",
        "测试类型",
        "前置条件",
        "测试步骤",
        "预期结果",
        "优先级",
        "执行结果"
    ]
    
    # Column widths
    COLUMN_WIDTHS = [12, 30, 12, 25, 40, 30, 10, 12]
    
    # Header styling
    HEADER_FILL = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    HEADER_FONT = Font(bold=True, color="FFFFFF", size=11)
    
    # Border style
    THIN_BORDER = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )
    
    def __init__(self):
        """Initialize the generator."""
        self.wb = None
        self.sheet = None
    
    def _setup_worksheet(self, sheet_name: str = "测试用例"):
        """Set up the worksheet with headers and formatting."""
        self.wb = Workbook()
        self.sheet = self.wb.active
        self.sheet.title = sheet_name
        
        # Set headers
        for col_idx, header in enumerate(self.HEADERS, 1):
            cell = self.sheet.cell(row=1, column=col_idx)
            cell.value = header
            cell.font = self.HEADER_FONT
            cell.fill = self.HEADER_FILL
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = self.THIN_BORDER
        
        # Set column widths
        for col_idx, width in enumerate(self.COLUMN_WIDTHS, 1):
            self.sheet.column_dimensions[get_column_letter(col_idx)].width = width
        
        # Freeze header row
        self.sheet.freeze_panes = "A2"
    
    def _format_data_row(self, row_num: int):
        """Apply formatting to a data row."""
        for col_idx in range(1, len(self.HEADERS) + 1):
            cell = self.sheet.cell(row=row_num, column=col_idx)
            cell.border = self.THIN_BORDER
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    
    def _generate_test_case_id(self, index: int) -> str:
        """Generate test case ID in format TC-001, TC-002, etc."""
        return f"TC-{index:03d}"
    
    def validate_test_case(self, test_case: Dict[str, str]) -> Tuple[bool, Optional[str]]:
        """
        Validate a test case dictionary.
        
        Args:
            test_case: Test case dictionary to validate
        
        Returns:
            Tuple of (is_valid, error_message)
        """
        required_fields = ['name', 'preconditions', 'steps', 'expected_result']
        missing_fields = [field for field in required_fields if not test_case.get(field)]
        
        if missing_fields:
            return False, f"Missing required fields: {', '.join(missing_fields)}"
        
        # Validate priority if provided
        if 'priority' in test_case and test_case['priority'] not in ['P0', 'P1', 'P2']:
            return False, f"Invalid priority: {test_case['priority']}. Must be P0, P1, or P2"
        
        return True, None
    
    def add_test_case(self, 
                     name: str,
                     preconditions: str,
                     steps: str,
                     expected_result: str,
                     priority: str = "P1",
                     test_type: str = "功能测试",
                     execution_result: Optional[str] = None,
                     validate: bool = True) -> int:
        """
        Add a single test case to the worksheet.
        
        Args:
            name: Test case name
            preconditions: Preconditions for test execution
            steps: Test steps (can be multi-line)
            expected_result: Expected result
            priority: Priority (P0, P1, P2)
            test_type: Test type (default: "功能测试")
            execution_result: Execution result (optional, defaults to empty)
            validate: Whether to validate the test case before adding (default: True)
        
        Returns:
            Row number where test case was added
        
        Raises:
            ValueError: If validation fails and validate=True
        """
        if validate:
            test_case_dict = {
                'name': name,
                'preconditions': preconditions,
                'steps': steps,
                'expected_result': expected_result,
                'priority': priority,
                'test_type': test_type
            }
            is_valid, error_msg = self.validate_test_case(test_case_dict)
            if not is_valid:
                raise ValueError(f"Invalid test case: {error_msg}")
        
        if self.sheet is None:
            self._setup_worksheet()
        
        # Get next row number
        row_num = self.sheet.max_row + 1
        
        # Generate test case ID
        test_case_id = self._generate_test_case_id(row_num - 1)
        
        # Add data
        data = [
            test_case_id,
            name,
            test_type,
            preconditions,
            steps,
            expected_result,
            priority,
            execution_result or ""
        ]
        
        for col_idx, value in enumerate(data, 1):
            cell = self.sheet.cell(row=row_num, column=col_idx)
            cell.value = value
        
        # Apply formatting
        self._format_data_row(row_num)
        
        return row_num
    
    def add_test_cases(self, test_cases: List[Dict[str, str]], validate: bool = True, skip_invalid: bool = False):
        """
        Add multiple test cases from a list of dictionaries.
        
        Args:
            test_cases: List of test case dictionaries with keys:
                - name (required)
                - preconditions (required)
                - steps (required)
                - expected_result (required)
                - priority (optional, default: "P1")
                - test_type (optional, default: "功能测试")
                - execution_result (optional)
            validate: Whether to validate test cases before adding (default: True)
            skip_invalid: If True, skip invalid test cases instead of raising error (default: False)
        
        Returns:
            Tuple of (added_count, skipped_count, errors)
        
        Raises:
            ValueError: If validation fails, validate=True, and skip_invalid=False
        """
        added_count = 0
        skipped_count = 0
        errors = []
        
        for i, test_case in enumerate(test_cases, 1):
            if validate:
                is_valid, error_msg = self.validate_test_case(test_case)
                if not is_valid:
                    error = f"Test case {i}: {error_msg}"
                    errors.append(error)
                    if skip_invalid:
                        skipped_count += 1
                        continue
                    else:
                        raise ValueError(error)
            
            try:
                self.add_test_case(
                    name=test_case.get("name", ""),
                    preconditions=test_case.get("preconditions", ""),
                    steps=test_case.get("steps", ""),
                    expected_result=test_case.get("expected_result", ""),
                    priority=test_case.get("priority", "P1"),
                    test_type=test_case.get("test_type", "功能测试"),
                    execution_result=test_case.get("execution_result"),
                    validate=False  # Already validated above
                )
                added_count += 1
            except Exception as e:
                error = f"Test case {i}: {str(e)}"
                errors.append(error)
                if skip_invalid:
                    skipped_count += 1
                    continue
                else:
                    raise
        
        return added_count, skipped_count, errors
    
    @staticmethod
    def load_from_file(file_path: str) -> List[Dict[str, str]]:
        """
        Load test cases from a JSON or YAML file.
        
        Args:
            file_path: Path to JSON or YAML file
        
        Returns:
            List of test case dictionaries
        
        Raises:
            FileNotFoundError: If file doesn't exist
            ValueError: If file format is invalid or unsupported
        """
        file_path = Path(file_path)
        if not file_path.exists():
            raise FileNotFoundError(f"File not found: {file_path}")
        
        suffix = file_path.suffix.lower()
        
        try:
            with open(file_path, 'r', encoding='utf-8') as f:
                if suffix in ['.json']:
                    data = json.load(f)
                elif suffix in ['.yaml', '.yml']:
                    if yaml is None:
                        raise ValueError("YAML support requires PyYAML. Install with: pip install pyyaml")
                    data = yaml.safe_load(f)
                else:
                    raise ValueError(f"Unsupported file format: {suffix}. Supported: .json, .yaml, .yml")
            
            if not isinstance(data, list):
                raise ValueError("File must contain a list of test cases")
            
            return data
        except json.JSONDecodeError as e:
            raise ValueError(f"Invalid JSON format: {e}")
        except yaml.YAMLError as e:
            raise ValueError(f"Invalid YAML format: {e}")
        except Exception as e:
            raise ValueError(f"Error reading file: {e}")
    
    def generate_excel(self, output_path: str, test_cases: Optional[List[Dict[str, str]]] = None, 
                      validate: bool = True, skip_invalid: bool = False):
        """
        Generate Excel file with test cases.
        
        Args:
            output_path: Path to save the Excel file
            test_cases: Optional list of test cases to add
            validate: Whether to validate test cases (default: True)
            skip_invalid: If True, skip invalid test cases instead of raising error (default: False)
        
        Returns:
            Path to generated Excel file
        
        Raises:
            ValueError: If validation fails and skip_invalid=False
            IOError: If file cannot be saved
        """
        if self.sheet is None:
            self._setup_worksheet()
        
        if test_cases:
            added, skipped, errors = self.add_test_cases(test_cases, validate=validate, skip_invalid=skip_invalid)
            if skipped > 0:
                print(f"Warning: {skipped} test case(s) were skipped due to validation errors")
            if errors and not skip_invalid:
                raise ValueError(f"Validation errors: {'; '.join(errors)}")
        
        try:
            output_path = Path(output_path)
            output_path.parent.mkdir(parents=True, exist_ok=True)
            self.wb.save(str(output_path))
            return str(output_path)
        except Exception as e:
            raise IOError(f"Failed to save Excel file: {e}")


def main():
    """Example usage of the test case generator."""
    generator = TestCaseGenerator()
    
    # Example test cases
    test_cases = [
        {
            "name": "用户登录功能-正常登录",
            "preconditions": "系统已启动，用户账号已注册",
            "steps": "1. 打开登录页面\n2. 输入正确的用户名和密码\n3. 点击登录按钮",
            "expected_result": "登录成功，跳转到主页",
            "priority": "P0"
        },
        {
            "name": "用户登录功能-错误密码",
            "preconditions": "系统已启动，用户账号已注册",
            "steps": "1. 打开登录页面\n2. 输入正确的用户名和错误的密码\n3. 点击登录按钮",
            "expected_result": "显示错误提示：密码错误",
            "priority": "P1"
        }
    ]
    
    generator.generate_excel("test_cases.xlsx", test_cases)
    print("Test case Excel file generated successfully!")


if __name__ == "__main__":
    main()
